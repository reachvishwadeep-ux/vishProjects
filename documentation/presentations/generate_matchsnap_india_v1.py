from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "MatchSnap_India_Concept_and_Execution_Plan_v1.pptx"
ASSETS = ROOT / "assets"

SLIDE_W = 13.333
SLIDE_H = 7.5

INK = "17142F"
MUTED = "6F6A82"
PURPLE = "6547ED"
VIOLET = "9A5CF4"
LILAC = "EEE9FF"
PALE = "F8F6FD"
WHITE = "FFFFFF"
MINT = "28B67D"
MINT_BG = "E3F7EF"
AMBER = "F0A34A"
AMBER_BG = "FFF2E1"
RED = "D85D68"
RED_BG = "FDEBED"
BLUE = "4275E8"
BLUE_BG = "EAF0FF"
LINE = "DDD8EC"
DARK_PURPLE = "2D205F"

FONT = "Aptos"
FONT_DISPLAY = "Aptos Display"


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def add_rect(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    fill: str,
    radius: bool = False,
    line: str | None = None,
    line_width: float = 1,
):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(
        shape_type, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(line_width)
    else:
        shape.line.fill.background()
    return shape


def add_line(
    slide,
    x1: float,
    y1: float,
    x2: float,
    y2: float,
    color: str = LINE,
    width: float = 1.5,
):
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1),
        Inches(y1),
        Inches(x2),
        Inches(y2),
    )
    line.line.color.rgb = rgb(color)
    line.line.width = Pt(width)
    return line


def add_text(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    size: float = 18,
    color: str = INK,
    bold: bool = False,
    font: str = FONT,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    valign: MSO_ANCHOR = MSO_ANCHOR.TOP,
    margin: float = 0.04,
    line_spacing: float = 1.0,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.line_spacing = line_spacing
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return box


def add_rich_text(
    slide,
    runs: list[tuple[str, bool, str]],
    x: float,
    y: float,
    w: float,
    h: float,
    size: float = 18,
    align: PP_ALIGN = PP_ALIGN.LEFT,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.04)
    frame.margin_right = Inches(0.04)
    frame.margin_top = Inches(0.04)
    frame.margin_bottom = Inches(0.04)
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    for text, bold, color in runs:
        run = paragraph.add_run()
        run.text = text
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = rgb(color)
    return box


def add_badge(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    fill: str = LILAC,
    color: str = PURPLE,
):
    add_rect(slide, x, y, w, 0.34, fill, radius=True)
    add_text(
        slide,
        text.upper(),
        x + 0.06,
        y + 0.03,
        w - 0.12,
        0.24,
        10,
        color,
        True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )


def add_title(slide, eyebrow: str, title: str, subtitle: str | None = None):
    add_text(slide, eyebrow.upper(), 0.72, 0.45, 5.7, 0.25, 10, PURPLE, True)
    add_text(slide, title, 0.72, 0.76, 11.8, 0.62, 29, INK, True, FONT_DISPLAY)
    if subtitle:
        add_text(slide, subtitle, 0.72, 1.36, 11.6, 0.4, 13.5, MUTED)


def add_footer(slide, number: int, dark: bool = False):
    color = "D7CEF9" if dark else "8A849E"
    add_text(
        slide,
        "MatchSnap India  •  Concept & execution plan  •  Version 1",
        0.72,
        7.16,
        5.5,
        0.18,
        8.5,
        color,
    )
    add_text(
        slide,
        f"{number:02d}",
        12.02,
        7.12,
        0.58,
        0.2,
        9,
        color,
        True,
        align=PP_ALIGN.RIGHT,
    )


def add_bullet_list(
    slide,
    items: list[str],
    x: float,
    y: float,
    w: float,
    h: float,
    size: float = 15,
    color: str = INK,
    bullet_color: str = PURPLE,
    gap: float = 0.55,
):
    for index, item in enumerate(items):
        item_y = y + index * gap
        add_rect(slide, x, item_y + 0.08, 0.1, 0.1, bullet_color, radius=True)
        add_text(slide, item, x + 0.22, item_y, w - 0.22, gap, size, color)


def add_card(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    label: str,
    title: str,
    body: str,
    accent: str = PURPLE,
    fill: str = WHITE,
):
    add_rect(slide, x, y, w, h, fill, radius=True, line=LINE)
    add_rect(slide, x, y, 0.08, h, accent, radius=True)
    add_text(slide, label.upper(), x + 0.26, y + 0.25, w - 0.5, 0.22, 9, accent, True)
    add_text(slide, title, x + 0.26, y + 0.62, w - 0.5, 0.5, 18, INK, True)
    add_text(slide, body, x + 0.26, y + 1.16, w - 0.5, h - 1.34, 12.5, MUTED)


def add_picture_cover(
    slide, path: Path, x: float, y: float, w: float, h: float
):
    with Image.open(path) as image:
        image_w, image_h = image.size
    image_ratio = image_w / image_h
    frame_ratio = w / h
    picture = slide.shapes.add_picture(
        str(path), Inches(x), Inches(y), Inches(w), Inches(h)
    )
    if image_ratio > frame_ratio:
        visible_ratio = frame_ratio / image_ratio
        crop = (1 - visible_ratio) / 2
        picture.crop_left = crop
        picture.crop_right = crop
    else:
        visible_ratio = image_ratio / frame_ratio
        crop = (1 - visible_ratio) / 2
        picture.crop_top = crop
        picture.crop_bottom = crop
    return picture


def add_phone(slide, path: Path, x: float, y: float, w: float, h: float):
    add_rect(slide, x - 0.07, y - 0.08, w + 0.14, h + 0.16, INK, radius=True)
    add_picture_cover(slide, path, x, y, w, h)


def blank_slide(prs: Presentation, fill: str = PALE):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background
    background.fill.solid()
    background.fill.fore_color.rgb = rgb(fill)
    return slide


def cover_slide(prs: Presentation):
    slide = blank_slide(prs, DARK_PURPLE)
    add_rect(slide, 9.0, 0.0, 4.1, 4.1, PURPLE, radius=True)
    add_rect(slide, 10.15, 3.25, 3.05, 4.1, VIOLET, radius=True)
    add_rect(slide, 9.42, 1.07, 2.05, 2.05, "BDA4FF", radius=True)
    add_text(slide, "MATCHSNAP", 0.75, 0.62, 2.2, 0.28, 11, "D9D0FF", True)
    add_badge(slide, "India-first • Version 1", 0.75, 1.28, 2.1, "463582", WHITE)
    add_text(
        slide,
        "A safer path from\nmissing to found",
        0.75,
        2.0,
        7.25,
        1.65,
        38,
        WHITE,
        True,
        FONT_DISPLAY,
    )
    add_text(
        slide,
        "Concept, product direction and execution plan for a partner-led India pilot.",
        0.78,
        4.0,
        6.5,
        0.72,
        18,
        "D9D0FF",
    )
    add_line(slide, 0.78, 5.12, 6.85, 5.12, "7562B5", 1)
    add_text(
        slide,
        "Remote face matching  •  Human review  •  Protected connection",
        0.78,
        5.35,
        6.65,
        0.38,
        13,
        WHITE,
        True,
    )
    add_text(
        slide,
        "Planning document — not a production-readiness or legal-compliance claim",
        0.78,
        6.52,
        6.8,
        0.28,
        10.5,
        "BEB4E1",
    )
    add_footer(slide, 1, dark=True)


def idea_slide(prs: Presentation):
    slide = blank_slide(prs)
    add_title(
        slide,
        "The concept",
        "Turn two disconnected photos into one reviewed lead",
        "MatchSnap searches a shared missing/found repository and gives authorized people a candidate to review.",
    )
    add_card(
        slide,
        0.72,
        2.12,
        3.42,
        2.45,
        "Missing repository",
        "Someone is being sought",
        "A family or authorized partner submits a verified case and suitable photograph.",
        PURPLE,
    )
    add_card(
        slide,
        9.19,
        2.12,
        3.42,
        2.45,
        "Found repository",
        "Someone has been found",
        "A trained finder or partner submits a sighting or found-person record.",
        MINT,
    )
    add_line(slide, 4.14, 3.34, 5.07, 3.34, PURPLE, 3)
    add_line(slide, 8.26, 3.34, 9.19, 3.34, MINT, 3)
    add_rect(slide, 5.07, 2.38, 3.19, 1.92, DARK_PURPLE, radius=True)
    add_text(
        slide,
        "REMOTE\nMATCHING",
        5.48,
        2.75,
        2.37,
        0.92,
        19,
        WHITE,
        True,
        align=PP_ALIGN.CENTER,
    )
    add_badge(slide, "Candidate, not identity proof", 4.91, 4.64, 3.5, AMBER_BG, AMBER)
    add_rect(slide, 0.72, 5.33, 11.89, 1.05, LILAC, radius=True)
    add_rich_text(
        slide,
        [
            ("Outcome: ", True, PURPLE),
            (
                "an authorized reviewer checks the candidate before any notification, handoff or connection.",
                False,
                INK,
            ),
        ],
        1.05,
        5.67,
        11.2,
        0.42,
        17,
        PP_ALIGN.CENTER,
    )
    add_footer(slide, 2)


def problem_slide(prs: Presentation):
    slide = blank_slide(prs)
    add_title(
        slide,
        "The opportunity",
        "The gap is not another public database",
        "The opportunity is a secure matching layer that fits authorized workflows and keeps people in control.",
    )
    cards = [
        (
            "01",
            "Time separates records",
            "A found-person image may arrive hours, days or months after the original missing report.",
            PURPLE,
        ),
        (
            "02",
            "Photos are compared manually",
            "People cannot reliably review growing photo collections across locations and time periods.",
            BLUE,
        ),
        (
            "03",
            "A lead needs safe handling",
            "A similarity score is useful only when verification, safeguarding and an official handoff follow.",
            MINT,
        ),
    ]
    for index, (label, title, body, accent) in enumerate(cards):
        x = 0.72 + index * 4.02
        add_card(slide, x, 2.13, 3.64, 3.04, label, title, body, accent)
    add_rect(slide, 0.72, 5.55, 11.89, 0.82, DARK_PURPLE, radius=True)
    add_text(
        slide,
        "Existing police and child-protection systems remain the backbone; MatchSnap should complement them.",
        1.05,
        5.78,
        11.2,
        0.32,
        15,
        WHITE,
        True,
        align=PP_ALIGN.CENTER,
    )
    add_footer(slide, 3)


def workflow_slide(prs: Presentation):
    slide = blank_slide(prs)
    add_title(
        slide,
        "The product",
        "A simple user flow with a controlled decision point",
        "The mobile experience stays simple; sensitive decisions move to authorized reviewers.",
    )
    add_phone(slide, ASSETS / "matchsnap-upload.png", 9.85, 1.8, 2.02, 4.49)
    steps = [
        ("1", "Choose intent", "Looking for someone or found someone"),
        ("2", "Submit photo", "Quality and consent checks"),
        ("3", "Search remotely", "Only the opposite repository"),
        ("4", "Persist candidates", "Immediate plus scheduled reconciliation"),
        ("5", "Human review", "Partner verifies the lead"),
        ("6", "Protected handoff", "No automatic contact disclosure"),
    ]
    for index, (number, title, body) in enumerate(steps):
        column = index % 2
        row = index // 2
        x = 0.72 + column * 4.25
        y = 1.95 + row * 1.48
        add_rect(slide, x, y, 0.54, 0.54, PURPLE if index < 4 else MINT, radius=True)
        add_text(
            slide,
            number,
            x + 0.04,
            y + 0.09,
            0.46,
            0.26,
            13,
            WHITE,
            True,
            align=PP_ALIGN.CENTER,
        )
        add_text(slide, title, x + 0.75, y - 0.01, 3.1, 0.32, 16, INK, True)
        add_text(slide, body, x + 0.75, y + 0.39, 3.18, 0.52, 11.5, MUTED)
        if row < 2:
            add_line(slide, x + 0.27, y + 0.58, x + 0.27, y + 1.33, LINE, 1.5)
    add_badge(slide, "Current prototype UI", 9.67, 6.47, 2.4, LILAC, PURPLE)
    add_footer(slide, 4)


def prototype_slide(prs: Presentation):
    slide = blank_slide(prs)
    add_title(
        slide,
        "Current foundation",
        "The remote matching proof of concept already works",
        "The prototype demonstrates the core technical loop; it does not yet provide production safeguards.",
    )
    nodes = [
        (0.8, 2.1, 2.2, "Flutter app", "Android + iOS"),
        (3.48, 2.1, 2.2, "FastAPI", "Remote recognition"),
        (6.16, 1.9, 2.35, "InsightFace", "Detect + embed"),
        (6.16, 2.95, 2.35, "PostgreSQL", "pgvector search"),
        (8.99, 2.95, 2.35, "MinIO", "Private photo store"),
    ]
    for x, y, w, title, body in nodes:
        fill = DARK_PURPLE if title == "FastAPI" else WHITE
        title_color = WHITE if title == "FastAPI" else INK
        body_color = "D8D0F5" if title == "FastAPI" else MUTED
        add_rect(slide, x, y, w, 1.0, fill, radius=True, line=LINE if fill == WHITE else None)
        add_text(slide, title, x + 0.18, y + 0.18, w - 0.36, 0.28, 15, title_color, True)
        add_text(slide, body, x + 0.18, y + 0.55, w - 0.36, 0.22, 10.5, body_color)
    add_line(slide, 3.0, 2.6, 3.48, 2.6, PURPLE, 2.5)
    add_line(slide, 5.68, 2.6, 6.16, 2.4, PURPLE, 2.5)
    add_line(slide, 5.68, 2.6, 6.16, 3.45, PURPLE, 2.5)
    add_line(slide, 8.51, 3.45, 8.99, 3.45, PURPLE, 2.5)
    proof = [
        ("Immediate", "Matching on upload"),
        ("Every 15 min", "Restart-safe reconciliation"),
        ("Opposite side", "Missing searches found, and vice versa"),
        ("Persisted", "Candidates remain reviewable"),
    ]
    for index, (value, label) in enumerate(proof):
        x = 0.8 + index * 3.0
        add_rect(slide, x, 4.8, 2.65, 1.25, WHITE, radius=True, line=LINE)
        add_text(slide, value, x + 0.2, 5.05, 2.25, 0.31, 18, PURPLE, True)
        add_text(slide, label, x + 0.2, 5.46, 2.25, 0.34, 10.5, MUTED)
    add_badge(slide, "Prototype — not production ready", 4.48, 6.39, 4.25, RED_BG, RED)
    add_footer(slide, 5)


def india_position_slide(prs: Presentation):
    slide = blank_slide(prs)
    add_title(
        slide,
        "India-first position",
        "Start as a partner tool, not a nationwide consumer launch",
        "A narrow launch creates a safer path to trust, evidence and eventual integration.",
    )
    add_rect(slide, 4.65, 2.2, 4.03, 2.1, DARK_PURPLE, radius=True)
    add_text(
        slide,
        "MATCHSNAP\nPILOT",
        5.28,
        2.73,
        2.77,
        0.88,
        24,
        WHITE,
        True,
        align=PP_ALIGN.CENTER,
    )
    systems = [
        (0.78, 1.95, 3.15, "Mission Vatsalya", "Child-protection framework"),
        (0.78, 4.45, 3.15, "DCPU / CWC", "District review and safeguarding"),
        (9.39, 1.95, 3.15, "TrackChild", "Existing official record workflow"),
        (9.39, 4.45, 3.15, "Child Helpline 1098", "Crisis and escalation channel"),
    ]
    for x, y, w, title, body in systems:
        add_rect(slide, x, y, w, 1.35, WHITE, radius=True, line=LINE)
        add_text(slide, title, x + 0.22, y + 0.26, w - 0.44, 0.33, 16, INK, True)
        add_text(slide, body, x + 0.22, y + 0.72, w - 0.44, 0.3, 11, MUTED)
    add_line(slide, 3.93, 2.63, 4.65, 2.78, PURPLE, 2)
    add_line(slide, 3.93, 5.12, 4.65, 3.78, PURPLE, 2)
    add_line(slide, 8.68, 2.78, 9.39, 2.63, PURPLE, 2)
    add_line(slide, 8.68, 3.78, 9.39, 5.12, PURPLE, 2)
    add_rect(slide, 4.15, 5.05, 5.03, 1.12, MINT_BG, radius=True)
    add_text(
        slide,
        "Recommended first shape\n1 population • 1 district • 1 accountable partner",
        4.45,
        5.32,
        4.43,
        0.62,
        15,
        MINT,
        True,
        align=PP_ALIGN.CENTER,
    )
    add_footer(slide, 6)


def guardrails_slide(prs: Presentation):
    slide = blank_slide(prs)
    add_title(
        slide,
        "Safety model",
        "Recognition proposes; people decide",
        "The service must reduce search effort without turning a sensitive database into an open surveillance tool.",
    )
    add_rect(slide, 0.72, 2.02, 5.78, 4.16, MINT_BG, radius=True)
    add_text(slide, "DESIGN FOR", 1.05, 2.34, 2.0, 0.26, 10, MINT, True)
    add_bullet_list(
        slide,
        [
            "Authorized, role-based access",
            "Verified cases and trained reviewers",
            "Human review before contact or escalation",
            "Private images, audit logs and deletion",
            "Plain-language consent and grievance handling",
        ],
        1.05,
        2.86,
        5.0,
        2.8,
        14.5,
        INK,
        MINT,
        0.58,
    )
    add_rect(slide, 6.82, 2.02, 5.79, 4.16, RED_BG, radius=True)
    add_text(slide, "NEVER ENABLE", 7.15, 2.34, 2.2, 0.26, 10, RED, True)
    add_bullet_list(
        slide,
        [
            "Public or unrestricted face search",
            "A score presented as confirmed identity",
            "Automatic disclosure of phone or location",
            "Direct unknown-adult-to-child connection",
            "Silent reuse of biometric data",
        ],
        7.15,
        2.86,
        5.0,
        2.8,
        14.5,
        INK,
        RED,
        0.58,
    )
    add_footer(slide, 7)


def legal_slide(prs: Presentation):
    slide = blank_slide(prs)
    add_title(
        slide,
        "India readiness",
        "Legal, safeguarding and licensing are product requirements",
        "Qualified Indian counsel and the pilot partner must approve the real-data workflow.",
    )
    items = [
        (
            "DPDP",
            "Notice, consent, rights, grievance, safeguards, breach response and current enforcement timeline.",
            PURPLE,
            LILAC,
        ),
        (
            "Children",
            "Guardian authority, child-data duties, safeguarding and Mission Vatsalya / JJ workflows.",
            BLUE,
            BLUE_BG,
        ),
        (
            "Data lifecycle",
            "Purpose limits, processor contracts, India-region hosting, retention, erasure and backup deletion.",
            MINT,
            MINT_BG,
        ),
        (
            "Model licence",
            "Resolve commercial rights for pretrained weights or select an acceptable alternative.",
            AMBER,
            AMBER_BG,
        ),
    ]
    for index, (label, body, accent, fill) in enumerate(items):
        x = 0.72 + (index % 2) * 6.08
        y = 2.06 + (index // 2) * 1.75
        add_rect(slide, x, y, 5.79, 1.43, fill, radius=True)
        add_rect(slide, x + 0.25, y + 0.28, 1.28, 0.36, accent, radius=True)
        add_text(
            slide,
            label.upper(),
            x + 0.32,
            y + 0.35,
            1.14,
            0.2,
            9,
            WHITE,
            True,
            align=PP_ALIGN.CENTER,
        )
        add_text(slide, body, x + 1.77, y + 0.24, 3.7, 0.85, 12.5, INK)
    add_rect(slide, 0.72, 5.73, 11.89, 0.62, DARK_PURPLE, radius=True)
    add_text(
        slide,
        "No real-person pilot until counsel, the partner and a named safety owner approve the controls.",
        1.05,
        5.91,
        11.2,
        0.27,
        14,
        WHITE,
        True,
        align=PP_ALIGN.CENTER,
    )
    add_footer(slide, 8)


def pilot_choice_slide(prs: Presentation):
    slide = blank_slide(prs)
    add_title(
        slide,
        "Phase 1 decision",
        "Choose the smallest responsible pilot",
        "Population determines consent, safeguarding, partner type and product workflow.",
    )
    choices = [
        (
            "A",
            "Children only",
            "Highest safeguarding burden; strongest fit with Mission Vatsalya, DCPU, CWC and 1098.",
            AMBER,
        ),
        (
            "B",
            "Adults only",
            "Simpler first consent model; requires police/NGO verification and vulnerable-adult safeguards.",
            BLUE,
        ),
        (
            "C",
            "Children + adults",
            "Broadest mission, but doubles workflow, policy and partner complexity for an initial pilot.",
            PURPLE,
        ),
    ]
    for index, (label, title, body, accent) in enumerate(choices):
        x = 0.72 + index * 4.02
        add_rect(slide, x, 2.1, 3.64, 2.9, WHITE, radius=True, line=LINE)
        add_rect(slide, x + 0.26, 2.4, 0.58, 0.58, accent, radius=True)
        add_text(
            slide,
            label,
            x + 0.33,
            2.54,
            0.44,
            0.22,
            13,
            WHITE,
            True,
            align=PP_ALIGN.CENTER,
        )
        add_text(slide, title, x + 0.26, 3.2, 3.1, 0.38, 19, INK, True)
        add_text(slide, body, x + 0.26, 3.78, 3.08, 0.92, 12, MUTED)
    add_rect(slide, 0.72, 5.45, 11.89, 0.91, LILAC, radius=True)
    add_rich_text(
        slide,
        [
            ("Recommendation: ", True, PURPLE),
            (
                "select one population, one city/district and one accountable institutional partner.",
                False,
                INK,
            ),
        ],
        1.05,
        5.72,
        11.2,
        0.36,
        15.5,
        PP_ALIGN.CENTER,
    )
    add_footer(slide, 9)


def roadmap_slide(prs: Presentation):
    slide = blank_slide(prs)
    add_title(
        slide,
        "Execution roadmap",
        "Nine phases, with evidence gates between them",
        "Do not scale engineering or outreach faster than safety, partner and evidence readiness.",
    )
    phases = [
        ("1", "Scope", "India charter"),
        ("2", "Validate", "Interviews + partner"),
        ("3", "Approve", "Legal + safeguarding"),
        ("4", "Harden", "Secure pilot MVP"),
        ("5", "Measure", "Accuracy + security"),
        ("6", "Pilot", "Closed real workflow"),
        ("7", "Model", "Operations + budget"),
        ("8", "Package", "Deck + data room"),
        ("9", "Fund", "Targeted outreach"),
    ]
    wave_titles = ["FOUNDATION", "BUILD & PROVE", "FUND & SCALE"]
    wave_colors = [PURPLE, BLUE, MINT]
    for wave, wave_title in enumerate(wave_titles):
        x = 0.72 + wave * 4.02
        w = 3.69
        color = wave_colors[wave]
        add_rect(slide, x, 1.9, w, 2.78, WHITE, radius=True, line=LINE)
        add_text(
            slide,
            wave_title,
            x + 0.24,
            2.14,
            w - 0.48,
            0.25,
            10,
            color,
            True,
        )
        add_line(slide, x + 0.24, 2.51, x + w - 0.24, 2.51, color, 2.5)
        for within, (number, title, body) in enumerate(phases[wave * 3 : wave * 3 + 3]):
            row_y = 2.75 + within * 0.61
            add_rect(slide, x + 0.24, row_y, 0.38, 0.38, color, radius=True)
            add_text(
                slide,
                number,
                x + 0.285,
                row_y + 0.095,
                0.29,
                0.16,
                9.5,
                WHITE,
                True,
                align=PP_ALIGN.CENTER,
            )
            add_text(slide, title, x + 0.78, row_y - 0.01, 0.9, 0.24, 11, INK, True)
            add_text(slide, body, x + 1.67, row_y - 0.01, 1.7, 0.42, 9.2, MUTED)
    gates = [
        ("GATE 1", "Partner confirms the workflow is useful and safe to explore."),
        ("GATE 2", "Counsel and partner approve real-data pilot controls."),
        ("GATE 3", "Pilot evidence supports a costed funding request."),
    ]
    for index, (label, body) in enumerate(gates):
        x = 0.72 + index * 4.02
        add_rect(slide, x, 5.15, 3.69, 1.02, WHITE, radius=True, line=LINE)
        add_text(
            slide,
            label,
            x + 0.22,
            5.36,
            0.85,
            0.2,
            9,
            wave_colors[index],
            True,
        )
        add_text(slide, body, x + 1.03, 5.27, 2.38, 0.58, 10.5, INK)
    add_footer(slide, 10)


def mvp_slide(prs: Presentation):
    slide = blank_slide(prs)
    add_title(
        slide,
        "Secure pilot MVP",
        "What must be added before real cases",
        "The next engineering milestone is controlled access and case handling—not more matching demos.",
    )
    groups = [
        (
            "IDENTITY & ACCESS",
            ["OTP login", "Partner membership", "Reviewer/admin roles", "Rate limits"],
            PURPLE,
            LILAC,
        ),
        (
            "CASE GOVERNANCE",
            ["Case verification", "Assignment + closure", "Audit trail", "Duplicate handling"],
            BLUE,
            BLUE_BG,
        ),
        (
            "PRIVACY & SAFETY",
            ["Consent records", "Retention + deletion", "Private images", "Protected handoff"],
            MINT,
            MINT_BG,
        ),
        (
            "INDIA EXPERIENCE",
            ["Low-bandwidth upload", "Retry + draft", "English + Hindi", "Launch-state language"],
            AMBER,
            AMBER_BG,
        ),
    ]
    for index, (title, items, accent, fill) in enumerate(groups):
        x = 0.72 + (index % 2) * 6.08
        y = 2.06 + (index // 2) * 2.05
        add_rect(slide, x, y, 5.79, 1.72, fill, radius=True)
        add_text(slide, title, x + 0.27, y + 0.27, 2.65, 0.24, 10, accent, True)
        for item_index, item in enumerate(items):
            item_x = x + 0.27 + (item_index % 2) * 2.72
            item_y = y + 0.77 + (item_index // 2) * 0.46
            add_rect(slide, item_x, item_y + 0.04, 0.09, 0.09, accent, radius=True)
            add_text(slide, item, item_x + 0.18, item_y, 2.35, 0.25, 11.5, INK)
    add_footer(slide, 11)


def validation_slide(prs: Presentation):
    slide = blank_slide(prs)
    add_title(
        slide,
        "Recognition evidence",
        "Measure failure modes, not just impressive matches",
        "Thresholds must be calibrated on consented, representative Indian pilot conditions.",
    )
    add_phone(slide, ASSETS / "matchsnap-potential-match.png", 0.88, 1.86, 1.88, 4.18)
    add_badge(slide, "Human review required", 0.77, 6.25, 2.15, AMBER_BG, AMBER)
    metrics = [
        ("TAR @ FAR", "Correct matches at a stated false-accept rate"),
        ("Rank recall", "Whether the right person appears in the review list"),
        ("False candidates", "Reviewer burden and harm potential per case"),
        ("Group performance", "Age, skin tone, gender and image-quality differences"),
        ("Quality rejection", "No face, multiple faces, blur, pose and occlusion"),
        ("Latency + cost", "End-to-end processing time and infrastructure cost"),
    ]
    for index, (title, body) in enumerate(metrics):
        column = index % 2
        row = index // 2
        x = 3.42 + column * 4.58
        y = 2.0 + row * 1.35
        add_rect(slide, x, y, 4.23, 1.03, WHITE, radius=True, line=LINE)
        add_text(slide, title, x + 0.24, y + 0.17, 1.56, 0.58, 12.5, PURPLE, True)
        add_text(slide, body, x + 1.93, y + 0.14, 2.05, 0.68, 10.2, MUTED)
    add_rect(slide, 3.42, 6.06, 8.81, 0.5, RED_BG, radius=True)
    add_text(
        slide,
        "Stop if any group shows unacceptable error, reviewer burden or safety risk.",
        3.75,
        6.2,
        8.15,
        0.21,
        12,
        RED,
        True,
        align=PP_ALIGN.CENTER,
    )
    add_footer(slide, 12)


def pilot_slide(prs: Presentation):
    slide = blank_slide(prs)
    add_title(
        slide,
        "Closed pilot",
        "Prove usefulness inside one supervised workflow",
        "The pilot is an operational study with a kill switch—not a public launch.",
    )
    bands = [
        (
            "SETUP",
            "One location\nOne partner\nNamed reviewers",
            PURPLE,
            LILAC,
        ),
        (
            "OPERATE",
            "Verified cases\nControlled candidates\nPartner handoff",
            BLUE,
            BLUE_BG,
        ),
        (
            "LEARN",
            "Weekly safety review\nSupport log\nParticipant feedback",
            MINT,
            MINT_BG,
        ),
        (
            "DECIDE",
            "Continue\nRevise\nOr stop",
            AMBER,
            AMBER_BG,
        ),
    ]
    for index, (title, body, accent, fill) in enumerate(bands):
        x = 0.72 + index * 3.01
        add_rect(slide, x, 2.25, 2.66, 2.7, fill, radius=True)
        add_text(
            slide,
            f"{index + 1:02d}",
            x + 0.26,
            2.55,
            0.55,
            0.25,
            11,
            accent,
            True,
        )
        add_text(slide, title, x + 0.26, 3.0, 2.1, 0.35, 17, INK, True)
        add_text(slide, body, x + 0.26, 3.55, 2.1, 0.95, 13, MUTED)
        if index < len(bands) - 1:
            add_line(slide, x + 2.66, 3.61, x + 3.01, 3.61, LINE, 2)
    add_rect(slide, 0.72, 5.41, 11.89, 0.95, DARK_PURPLE, radius=True)
    add_text(
        slide,
        "Publish impact claims only when the partner verifies the outcome and approves communication.",
        1.12,
        5.72,
        11.1,
        0.3,
        14,
        WHITE,
        True,
        align=PP_ALIGN.CENTER,
    )
    add_footer(slide, 13)


def scorecard_slide(prs: Presentation):
    slide = blank_slide(prs)
    add_title(
        slide,
        "Pilot scorecard",
        "Success combines product, safety and operations",
        "A single recognition score cannot demonstrate funder readiness.",
    )
    columns = [
        (
            "PRODUCT",
            [
                "Verified cases completed",
                "Valid-photo upload rate",
                "Processing + review time",
                "Candidate-list usefulness",
            ],
            PURPLE,
        ),
        (
            "SAFETY",
            [
                "False candidates per case",
                "User understanding",
                "Deletion + grievance SLA",
                "Privacy/safeguarding incidents",
            ],
            RED,
        ),
        (
            "OPERATIONS",
            [
                "Reviewer time per case",
                "Notification + handoff rate",
                "Partner satisfaction",
                "Cost per processed case",
            ],
            MINT,
        ),
    ]
    for index, (title, items, accent) in enumerate(columns):
        x = 0.72 + index * 4.02
        add_rect(slide, x, 2.08, 3.64, 3.58, WHITE, radius=True, line=LINE)
        add_rect(slide, x, 2.08, 3.64, 0.68, accent, radius=True)
        add_text(
            slide,
            title,
            x + 0.26,
            2.31,
            3.1,
            0.23,
            11,
            WHITE,
            True,
            align=PP_ALIGN.CENTER,
        )
        add_bullet_list(
            slide,
            items,
            x + 0.28,
            3.03,
            3.04,
            2.2,
            12.5,
            INK,
            accent,
            0.57,
        )
    add_badge(slide, "Pre-agree thresholds and stop conditions", 4.13, 6.04, 5.08, LILAC, PURPLE)
    add_footer(slide, 14)


def operating_slide(prs: Presentation):
    slide = blank_slide(prs)
    add_title(
        slide,
        "Operating model",
        "Funders will ask who is accountable after the demo",
        "The organization, partner and technology roles must be explicit before money is requested.",
    )
    add_rect(slide, 0.72, 2.0, 7.2, 4.26, WHITE, radius=True, line=LINE)
    roles = [
        ("MATCHSNAP", "Product, security, cloud, support and incident response", PURPLE),
        ("PILOT PARTNER", "Case verification, trained review, handoff and safeguarding", BLUE),
        ("ADVISORS", "Indian privacy, child protection, ethics and recognition quality", MINT),
    ]
    for index, (title, body, accent) in enumerate(roles):
        y = 2.38 + index * 1.15
        add_rect(slide, 1.05, y, 1.58, 0.35, accent, radius=True)
        add_text(
            slide,
            title,
            1.14,
            y + 0.07,
            1.4,
            0.18,
            8.5,
            WHITE,
            True,
            align=PP_ALIGN.CENTER,
        )
        add_text(slide, body, 2.91, y - 0.03, 4.48, 0.62, 12.5, INK)
    add_line(slide, 1.05, 5.55, 7.53, 5.55, LINE, 1)
    add_text(slide, "Entity options to evaluate", 1.05, 5.72, 1.88, 0.2, 9.5, PURPLE, True)
    add_text(
        slide,
        "Social enterprise  •  Section 8  •  NGO/government partner",
        3.05,
        5.69,
        4.45,
        0.3,
        9.8,
        MUTED,
    )
    add_rect(slide, 8.31, 2.0, 4.3, 4.26, DARK_PURPLE, radius=True)
    add_text(slide, "12-MONTH USE OF FUNDS", 8.65, 2.39, 3.62, 0.25, 10, "CFC5F3", True)
    fund_items = [
        ("01", "Legal, privacy and safeguarding"),
        ("02", "Secure product + cloud"),
        ("03", "Evaluation and security review"),
        ("04", "Partner pilot operations"),
        ("05", "Core team and support"),
    ]
    for index, (number, item) in enumerate(fund_items):
        y = 2.92 + index * 0.58
        add_text(slide, number, 8.65, y, 0.34, 0.2, 9.5, VIOLET, True)
        add_text(slide, item, 9.12, y - 0.02, 2.96, 0.26, 11.5, WHITE)
    add_footer(slide, 15)


def funding_slide(prs: Presentation):
    slide = blank_slide(prs)
    add_title(
        slide,
        "Funder readiness",
        "Evidence first, then targeted outreach",
        "The same story is packaged differently for government, CSR, grants, incubators and investors.",
    )
    add_rect(slide, 0.72, 2.01, 5.42, 4.23, WHITE, radius=True, line=LINE)
    add_text(slide, "THE PACKAGE", 1.05, 2.34, 1.7, 0.25, 10, PURPLE, True)
    add_bullet_list(
        slide,
        [
            "Two-minute working demo",
            "Partner letter and pilot workflow",
            "Measured recognition and safety results",
            "Privacy, safeguarding and licence path",
            "12-month budget and exact funding ask",
            "Roadmap, team plan and risk register",
        ],
        1.05,
        2.85,
        4.65,
        3.0,
        13.5,
        INK,
        PURPLE,
        0.52,
    )
    add_rect(slide, 6.52, 2.01, 6.09, 4.23, LILAC, radius=True)
    add_text(slide, "INDIA OUTREACH CHANNELS", 6.86, 2.34, 2.65, 0.25, 10, PURPLE, True)
    channels = [
        ("DPIIT / Startup India", "Recognition, incubators and suitable schemes"),
        ("IndiaAI", "Current startup financing and acceleration opportunities"),
        ("CSR + foundations", "Child protection, public safety and social impact"),
        ("Government innovation", "Relevant central and state programs"),
        ("Impact capital", "Mission-aligned angels, funds and accelerators"),
    ]
    for index, (title, body) in enumerate(channels):
        y = 2.81 + index * 0.67
        add_text(slide, title, 6.86, y, 2.0, 0.25, 11.5, INK, True)
        add_text(slide, body, 8.93, y - 0.01, 3.1, 0.46, 9.8, MUTED)
    add_footer(slide, 16)


def next_steps_slide(prs: Presentation):
    slide = blank_slide(prs)
    add_title(
        slide,
        "Next 30 days",
        "Turn the concept into an India project charter",
        "These actions should happen before more real-data collection or nationwide planning.",
    )
    weeks = [
        (
            "WEEK 1",
            "Decide",
            "Population\nState + district\nPartner type",
            PURPLE,
            LILAC,
        ),
        (
            "WEEK 2",
            "Validate",
            "Interview guide\n4–6 stakeholder calls\nWorkflow map",
            BLUE,
            BLUE_BG,
        ),
        (
            "WEEK 3",
            "Partner",
            "Pilot owner\nWritten interest\nSafeguarding review",
            MINT,
            MINT_BG,
        ),
        (
            "WEEK 4",
            "Scope",
            "Legal brief\nMVP requirements\nPilot budget",
            AMBER,
            AMBER_BG,
        ),
    ]
    for index, (week, title, body, accent, fill) in enumerate(weeks):
        x = 0.72 + index * 3.01
        add_rect(slide, x, 2.13, 2.66, 3.17, fill, radius=True)
        add_text(slide, week, x + 0.24, 2.45, 1.15, 0.22, 9.5, accent, True)
        add_text(slide, title, x + 0.24, 2.93, 2.12, 0.4, 19, INK, True)
        add_text(slide, body, x + 0.24, 3.62, 2.1, 1.02, 13, MUTED)
    add_rect(slide, 0.72, 5.7, 11.89, 0.64, DARK_PURPLE, radius=True)
    add_text(
        slide,
        "First decision: missing children, missing adults, or both?",
        1.05,
        5.89,
        11.2,
        0.27,
        15,
        WHITE,
        True,
        align=PP_ALIGN.CENTER,
    )
    add_footer(slide, 17)


def sources_slide(prs: Presentation):
    slide = blank_slide(prs, DARK_PURPLE)
    add_text(slide, "OFFICIAL STARTING REFERENCES", 0.75, 0.58, 4.5, 0.28, 11, "CFC5F3", True)
    add_text(slide, "India sources used for Version 1", 0.75, 1.04, 8.4, 0.65, 29, WHITE, True, FONT_DISPLAY)
    sources = [
        (
            "Digital Personal Data Protection Act, 2023",
            "https://www.indiacode.nic.in/handle/123456789/22037?view_type=browse",
        ),
        (
            "Digital Personal Data Protection Rules, 2025",
            "https://www.meity.gov.in/documents/act-and-policies/digital-personal-data-protection-rules-2025-gDOxUjMtQWa?hl=en-US",
        ),
        ("Mission Vatsalya", "https://www.wcd.gov.in/child/mission_vatsalya"),
        (
            "Child Helpline 1098",
            "https://www.wcd.gov.in/offerings/mission-vatsalya-child-helpline-1098",
        ),
        ("TrackChild", "https://trackthemissingchild.gov.in/trackchild/"),
        (
            "Startup India Seed Fund Scheme",
            "https://seedfund.startupindia.gov.in/",
        ),
        (
            "DPIIT startup recognition",
            "https://www.startupindia.gov.in/content/sih/en/startupgov/startup_recognition_page.html",
        ),
        (
            "IndiaAI startup financing",
            "https://indiaai.gov.in/hub/indiaai-startup-financing",
        ),
    ]
    for index, (title, url) in enumerate(sources):
        column = index % 2
        row = index // 2
        x = 0.75 + column * 6.15
        y = 2.05 + row * 1.06
        add_text(slide, title, x, y, 5.38, 0.26, 12.5, WHITE, True)
        add_text(slide, url, x, y + 0.36, 5.38, 0.38, 8.8, "BEB4E1")
        link = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(x),
            Inches(y + 0.34),
            Inches(5.38),
            Inches(0.42),
        )
        link.fill.background()
        link.line.fill.background()
        link.click_action.hyperlink.address = url
    add_rect(slide, 0.75, 6.41, 11.83, 0.36, "463582", radius=True)
    add_text(
        slide,
        "Confirm current law, scheme eligibility and program deadlines before acting.",
        1.05,
        6.5,
        11.22,
        0.18,
        10,
        "DED8F5",
        align=PP_ALIGN.CENTER,
    )
    add_footer(slide, 18, dark=True)


def set_properties(prs: Presentation):
    prs.core_properties.title = "MatchSnap India Concept and Execution Plan — Version 1"
    prs.core_properties.subject = (
        "India-first product concept, safety model, pilot roadmap and funder readiness"
    )
    prs.core_properties.author = "MatchSnap"
    prs.core_properties.keywords = (
        "MatchSnap, India, missing persons, face recognition, pilot, fundraising"
    )
    prs.core_properties.comments = (
        "Planning document. Similarity candidates require human review."
    )


def validate(prs: Presentation):
    if len(prs.slides) != 18:
        raise ValueError(f"Expected 18 slides, found {len(prs.slides)}")
    slide_width = prs.slide_width
    slide_height = prs.slide_height
    errors: list[str] = []
    for slide_number, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            if shape.left < 0 or shape.top < 0:
                errors.append(f"Slide {slide_number}: shape starts outside slide")
            if shape.left + shape.width > slide_width:
                errors.append(f"Slide {slide_number}: shape exceeds slide width")
            if shape.top + shape.height > slide_height:
                errors.append(f"Slide {slide_number}: shape exceeds slide height")
    if errors:
        raise ValueError("\n".join(errors))


def build_deck():
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    set_properties(prs)
    cover_slide(prs)
    idea_slide(prs)
    problem_slide(prs)
    workflow_slide(prs)
    prototype_slide(prs)
    india_position_slide(prs)
    guardrails_slide(prs)
    legal_slide(prs)
    pilot_choice_slide(prs)
    roadmap_slide(prs)
    mvp_slide(prs)
    validation_slide(prs)
    pilot_slide(prs)
    scorecard_slide(prs)
    operating_slide(prs)
    funding_slide(prs)
    next_steps_slide(prs)
    sources_slide(prs)
    validate(prs)
    prs.save(OUTPUT)
    print(f"Created {OUTPUT}")


if __name__ == "__main__":
    build_deck()
